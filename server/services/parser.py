import pymupdf  # PyMuPDF (formerly fitz)
import docx
from pptx import Presentation
import re
from pathlib import Path
from ..models import FileType

def clean_text(text: str) -> str:
    """Normalize whitespace and remove junk characters."""
    if not text:
        return ""
    # Replace multiple newlines with a single newline
    text = re.sub(r'\n+', '\n', text)
    # Replace multiple spaces with a single space
    text = re.sub(r' +', ' ', text)
    # Remove null bytes and other common control characters
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', text)
    return text.strip()

def parse_pdf(file_path: str) -> str:
    """Extract text from a PDF file."""
    text = ""
    try:
        doc = pymupdf.open(file_path)
        for page in doc:
            text += page.get_text("text") + "\n"
        doc.close()
    except Exception as e:
        raise Exception(f"Failed to parse PDF: {str(e)}")
    return clean_text(text)

def parse_docx(file_path: str) -> str:
    """Extract text from a DOCX file including tables and paragraphs."""
    text_pieces = []
    try:
        doc = docx.Document(file_path)
        # Attempt to read in document order via body elements
        if hasattr(doc, "element") and hasattr(doc.element, "body"):
            for element in doc.element.body:
                if element.tag.endswith("p"):
                    p = docx.text.paragraph.Paragraph(element, doc)
                    if p.text.strip():
                        text_pieces.append(p.text)
                elif element.tag.endswith("tbl"):
                    tbl = docx.table.Table(element, doc)
                    for row in tbl.rows:
                        row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        unique_cells = []
                        for c in row_cells:
                            if not unique_cells or c != unique_cells[-1]:
                                unique_cells.append(c)
                        if unique_cells:
                            text_pieces.append(" | ".join(unique_cells))
        if not text_pieces:
            for para in doc.paragraphs:
                if para.text.strip():
                    text_pieces.append(para.text)
            for tbl in doc.tables:
                for row in tbl.rows:
                    row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if row_cells:
                        text_pieces.append(" | ".join(row_cells))
    except Exception as e:
        raise Exception(f"Failed to parse DOCX: {str(e)}")
    return clean_text("\n".join(text_pieces))

def parse_pptx(file_path: str) -> str:
    """Extract text from a PPTX file."""
    text = ""
    try:
        prs = Presentation(file_path)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
    except Exception as e:
        raise Exception(f"Failed to parse PPTX: {str(e)}")
    return clean_text(text)

def parse_file(file_path: str, file_type: FileType) -> str:
    """Dispatcher to parse document based on file type."""
    if file_type == FileType.pdf:
        return parse_pdf(file_path)
    elif file_type == FileType.docx:
        return parse_docx(file_path)
    elif file_type == FileType.pptx:
        return parse_pptx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")
